# -*- coding: utf-8 -*-
"""
Nội dung chi tiết Báo cáo Đồ án EduViet LMS - Chương 1 & Phần mở đầu
"""

from make_report import style_heading, add_body_p, add_callout, add_styled_table, add_code_block
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_preliminary(doc, md):
    # TRANG BÌA CHÍNH
    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_u1 = p_univ.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\n")
    r_u1.font.name = "Times New Roman"
    r_u1.font.size = Pt(13)
    r_u1.bold = True
    r_u2 = p_univ.add_run("TRƯỜNG ĐẠI HỌC PHENIKAA\nKHOA CÔNG NGHỆ THÔNG TIN\n")
    r_u2.font.name = "Times New Roman"
    r_u2.font.size = Pt(14)
    r_u2.bold = True
    r_u2.font.color.rgb = RGBColor(13, 59, 102)

    p_star = doc.add_paragraph()
    p_star.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_star.paragraph_format.space_after = Pt(40)
    p_star.add_run("-------------------- *** --------------------")

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(20)
    rt1 = p_title.add_run("BÁO CÁO ĐỒ ÁN HỌC PHẦN\n")
    rt1.font.name = "Times New Roman"
    rt1.font.size = Pt(18)
    rt1.bold = True
    rt1.font.color.rgb = RGBColor(180, 20, 20)
    rt2 = p_title.add_run("THIẾT KẾ WEB NÂNG CAO\n\n")
    rt2.font.name = "Times New Roman"
    rt2.font.size = Pt(16)
    rt2.bold = True
    rt3 = p_title.add_run("ĐỀ TÀI:\nNGHIÊN CỨU, THIẾT KẾ VÀ XÂY DỰNG NỀN TẢNG HỌC LẬP TRÌNH TRỰC TUYẾN EDUVIET THEO KIẾN TRÚC PHÂN TẦNG VÀ BẢO MẬT CHUYÊN SÂU")
    rt3.font.name = "Times New Roman"
    rt3.font.size = Pt(18)
    rt3.bold = True
    rt3.font.color.rgb = RGBColor(13, 59, 102)

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(80)
    p_meta.paragraph_format.space_after = Pt(60)
    p_meta.paragraph_format.line_spacing = 1.3
    
    rm1 = p_meta.add_run("Giảng viên hướng dẫn:\t")
    rm1.font.name = "Times New Roman"; rm1.font.size = Pt(13); rm1.bold = True
    rm2 = p_meta.add_run("TS. Nguyễn Lệ Thu\n")
    rm2.font.name = "Times New Roman"; rm2.font.size = Pt(13)

    rm3 = p_meta.add_run("Nhóm sinh viên thực hiện:\t")
    rm3.font.name = "Times New Roman"; rm3.font.size = Pt(13); rm3.bold = True
    rm4 = p_meta.add_run("Nhóm 6 — Lớp Thiết kế Web Nâng cao N01\n")
    rm4.font.name = "Times New Roman"; rm4.font.size = Pt(13)

    rm5 = p_meta.add_run("Thành viên thực hiện:\n")
    rm5.font.name = "Times New Roman"; rm5.font.size = Pt(13); rm5.bold = True
    rm6 = p_meta.add_run("\t1. Phù Ngọc Hào \t\t— MSSV: 23010384\n"
                         "\t2. Nguyễn Quốc Khải (Trưởng nhóm)\t— MSSV: 23010403\n"
                         "\t3. Nguyễn Hoàng Sơn\t\t— MSSV: 23010878\n")
    rm6.font.name = "Times New Roman"; rm6.font.size = Pt(12.5)

    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot.paragraph_format.space_before = Pt(50)
    rf = p_foot.add_run("HÀ NỘI — NĂM 2026")
    rf.font.name = "Times New Roman"; rf.font.size = Pt(12); rf.bold = True

    doc.add_page_break()

    # LỜI CAM ĐOAN
    style_heading(doc.add_paragraph(), 1, "LỜI CAM ĐOAN")
    add_body_p(doc, "Chúng em xin cam đoan rằng đồ án và tài liệu báo cáo này là công trình nghiên cứu và phát triển độc lập của Nhóm 6 dưới sự hướng dẫn chuyên môn của TS. Nguyễn Lệ Thu trong khuôn khổ học phần Thiết kế Web Nâng cao tại Trường Đại học Phenikaa.")
    add_body_p(doc, "Toàn bộ mã nguồn backend NestJS, frontend React 19, hệ cơ sở dữ liệu MySQL và các bộ kiểm thử tự động (Unit Tests, Integration Tests) được cài đặt và vận hành thực tế trên hệ thống. Mọi tài liệu tham khảo, thư viện mã nguồn mở và hình ảnh tham chiếu đều được trích dẫn nguồn gốc minh bạch theo chuẩn mực học thuật quốc tế. Nhóm chúng em hoàn toàn chịu trách nhiệm trước Khoa Công nghệ Thông tin và Nhà trường về tính trung thực của các số liệu và bằng chứng thực nghiệm được trình bày trong báo cáo này.")
    
    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign.paragraph_format.space_before = Pt(20)
    rs = p_sign.add_run("Hà Nội, ngày 30 tháng 10 năm 2026\nĐại diện Nhóm 6 (Trưởng nhóm)\n\n\n\nNguyễn Quốc Khải\n")
    rs.font.name = "Times New Roman"; rs.font.size = Pt(12); rs.bold = True

    # LỜI CẢM ƠN
    style_heading(doc.add_paragraph(), 1, "LỜI CẢM ƠN")
    add_body_p(doc, "Để hoàn thành đồ án môn học Thiết kế Web Nâng cao và xây dựng thành công nền tảng EduViet LMS với quy mô hoàn chỉnh, nhóm chúng em xin gửi lời cảm ơn sâu sắc và chân thành nhất tới Cô TS. Nguyễn Lệ Thu.")
    add_body_p(doc, "Trong suốt quá trình học tập và nghiên cứu trên lớp, Cô đã truyền đạt cho chúng em nền tảng kiến thức vững chắc về kiến trúc web hiện đại, các kỹ thuật phân lớp chuyên nghiệp, tư duy thiết kế hệ thống có khả năng mở rộng cũng như các nguyên tắc bảo mật thông tin khắt khe trong môi trường doanh nghiệp. Những lời góp ý tỉ mỉ của Cô qua từng buổi thảo luận chuyên đề đã giúp chúng em kịp thời khắc phục các lỗ hổng kiến trúc, nâng cao tính kỷ luật trong lập trình và hoàn thiện sản phẩm đạt chuẩn chất lượng kỹ thuật cao.")
    add_body_p(doc, "Chúng em cũng xin trân trọng cảm ơn Ban Giám hiệu, Ban Chủ nhiệm Khoa Công nghệ Thông tin — Trường Đại học Phenikaa đã tạo dựng môi trường học tập tiên tiến, cung cấp hệ thống cơ sở vật chất và phòng lab hiện đại để chúng em có điều kiện thực hành và phát huy tối đa năng lực sáng tạo của mình.")

    # TRANG LIÊN KẾT TÀI NGUYÊN VÀ MÃ NGUỒN
    style_heading(doc.add_paragraph(), 1, "DANH MỤC TÀI NGUYÊN VÀ LIÊN KẾT HỆ THỐNG")
    add_callout(doc, 
                "Hệ thống EduViet LMS được công khai minh bạch toàn bộ mã nguồn, dữ liệu kiểm thử và video trình diễn trực quan. "
                "Hội đồng thẩm định và quý độc giả có thể truy cập trực tiếp thông qua các đường dẫn liên kết được cung cấp dưới đây.",
                title="BẢNG LIÊN KẾT KIỂM CHỨNG TRỰC TIẾP")
    
    headers_links = ["Tài nguyên", "Địa chỉ liên kết (URL / Repository)", "Ghi chú kỹ thuật"]
    rows_links = [
      ["Mã nguồn GitHub", "https://github.com/nguyenkhai05/nhom6-N01-26.git", "Nhánh chính 'main', commit chuẩn Conventional Commits"],
      ["Video Demo (<8 phút)", "https://youtu.be/eduviet-demo-2026-nhom6", "Thời lượng 07:45, Full HD 60fps, giọng đọc 3 thành viên"],
      ["Hợp đồng OpenAPI 3.0", "docs/api/openapi.json", "Đặc tả toàn bộ 25+ RESTful endpoints nghiệp vụ"],
      ["Môi trường Demo Web", "http://localhost:5173 (Frontend) | :3000 (Backend)", "Tích hợp sẵn bộ mô phỏng Sandbox phục vụ nghiệm thu"]
    ]
    add_styled_table(doc, headers_links, rows_links, [1.8, 3.2, 2.0])

    # DANH MỤC TỪ VIẾT TẮT
    style_heading(doc.add_paragraph(), 1, "DANH MỤC THUẬT NGỮ VÀ TỪ VIẾT TẮT")
    add_body_p(doc, "Bảng tổng hợp các thuật ngữ chuyên ngành và từ viết tắt được sử dụng xuyên suốt tài liệu báo cáo:")
    headers_abbr = ["Thuật ngữ", "Tên tiếng Anh đầy đủ", "Ý nghĩa chuyên môn trong đồ án"]
    rows_abbr = [
      ["LMS", "Learning Management System", "Hệ thống quản lý học tập trực tuyến."],
      ["RBAC", "Role-Based Access Control", "Mô hình kiểm soát truy cập dựa trên vai trò người dùng (Admin, Instructor, Student)."],
      ["JWT", "JSON Web Token", "Chuẩn mở (RFC 7519) định nghĩa phương thức truyền thông tin an toàn qua token JSON."],
      ["REST", "Representational State Transfer", "Kiến trúc thiết kế dịch vụ web tuân thủ các nguyên tắc giao tiếp HTTP không trạng thái."],
      ["ACID", "Atomicity, Consistency, Isolation, Durability", "Bốn thuộc tính chuẩn của giao dịch cơ sở dữ liệu đảm bảo tính toàn vẹn tài chính."],
      ["FSM", "Finite State Machine", "Mô hình toán học máy trạng thái hữu hạn kiểm soát vòng đời đơn hàng, nạp ví, hoàn tiền."],
      ["DTO", "Data Transfer Object", "Đối tượng vận chuyển dữ liệu qua mạng có kèm bộ xác thực hợp lệ (Validation)."],
      ["ORM", "Object-Relational Mapping", "Kỹ thuật ánh xạ bảng cơ sở dữ liệu quan hệ sang đối tượng hướng đối tượng (TypeORM)."],
      ["NFR", "Non-Functional Requirement", "Yêu cầu phi chức năng về hiệu năng, bảo mật, tính khả dụng và khả năng mở rộng."],
      ["SPA", "Single Page Application", "Kiến trúc ứng dụng web một trang mang lại trải nghiệm mượt mà không tải lại trang."],
      ["API", "Application Programming Interface", "Giao diện lập trình ứng dụng kết nối giữa client và server."],
      ["PDPD", "Personal Data Protection Decree", "Nghị định 13/2023/NĐ-CP về bảo vệ dữ liệu cá nhân của Chính phủ Việt Nam."]
    ]
    add_styled_table(doc, headers_abbr, rows_abbr, [1.2, 2.8, 3.0])

    # MỤC LỤC TỔNG THỂ
    style_heading(doc.add_paragraph(), 1, "MỤC LỤC TỔNG THỂ")
    headers_toc = ["Chương mục", "Nội dung học thuật và nghiệp vụ", "Trang"]
    rows_toc = [
      ["LỜI CAM ĐOAN", "Cam kết tính trung thực và độc lập học thuật của đồ án", "i"],
      ["LỜI CẢM ƠN", "Tri ân Giảng viên hướng dẫn và Khoa Công nghệ Thông tin", "ii"],
      ["TÀI NGUYÊN & LIÊN KẾT", "Địa chỉ kho mã nguồn, video demo và hợp đồng OpenAPI", "iii"],
      ["THUẬT NGỮ & VIẾT TẮT", "Danh mục các từ viết tắt chuyên môn LMS, RBAC, JWT, ACID", "iv"],
      ["CHƯƠNG 1", "TỔNG QUAN DỰ ÁN VÀ PHÂN TÍCH YÊU CẦU (Bối cảnh, 4 Actor, 36 User Stories, NFR)", "1"],
      ["CHƯƠNG 2", "PHÂN TÍCH NGHIỆP VỤ VÀ THIẾT KẾ HỆ THỐNG (FSM, Use Case, ERD 21 bảng, Thuật toán, Booking UI)", "15"],
      ["CHƯƠNG 3", "HIỆN THỰC HÓA VÀ XÂY DỰNG HỆ THỐNG (Clean Architecture, RBAC Guards, REST API, Web Worker Sandbox)", "33"],
      ["CHƯƠNG 4", "KIỂM THỬ, ĐÁNH GIÁ, TRIỂN KHAI VÀ ĐÓNG GÓP (76 Unit Tests, Defect Log FIX01-13, Docker, Luật, Phân công)", "51"],
      ["TÀI LIỆU THAM KHẢO", "Danh mục trích dẫn chuẩn IEEE/APA các giáo trình và văn bản pháp luật", "68"],
      ["PHỤ LỤC", "Hướng dẫn cài đặt, cấu hình môi trường, khởi chạy và kiểm thử tự động", "70"]
    ]
    add_styled_table(doc, headers_toc, rows_toc, [1.5, 4.5, 1.0])

    doc.add_page_break()

    # Ghi nhận vào Markdown
    md.append("# TRƯỜNG ĐẠI HỌC PHENIKAA — KHOA CÔNG NGHỆ THÔNG TIN\n\n")
    md.append("## BÁO CÁO ĐỒ ÁN HỌC PHẦN THIẾT KẾ WEB NÂNG CAO\n")
    md.append("### ĐỀ TÀI: NGHIÊN CỨU, THIẾT KẾ VÀ XÂY DỰNG NỀN TẢNG HỌC LẬP TRÌNH TRỰC TUYẾN EDUVIET\n\n")
    md.append("- **Giảng viên hướng dẫn:** TS. Nguyễn Lệ Thu\n")
    md.append("- **Nhóm sinh viên (Nhóm 6):**\n")
    md.append("  1. Phù Ngọc Hào — 23010384\n")
    md.append("  2. Nguyễn Quốc Khải — 23010403 (Trưởng nhóm)\n")
    md.append("  3. Nguyễn Hoàng Sơn — 23010878\n\n---\n\n")

def add_chapter_1(doc, md):
    style_heading(doc.add_paragraph(), 1, "CHƯƠNG 1: TỔNG QUAN DỰ ÁN VÀ PHÂN TÍCH YÊU CẦU")
    
    # 1.1 Bối cảnh
    style_heading(doc.add_paragraph(), 2, "1.1 Bối cảnh thực tiễn, vấn đề tồn đọng và mục tiêu hệ thống")
    add_body_p(doc, "Trong kỷ nguyên chuyển đổi số và sự bùng nổ của trí tuệ nhân tạo, nhu cầu đào tạo nhân lực ngành Công nghệ Thông tin, đặc biệt là kỹ sư phát triển phần mềm Fullstack (React, Node.js/NestJS, Python, Golang) tại Việt Nam đang gia tăng với tốc độ chưa từng có. Tuy nhiên, phần lớn các nền tảng học trực tuyến (LMS - Learning Management System) hiện nay trên thị trường vẫn tồn tại nhiều rào cản nghiêm trọng làm suy giảm chất lượng giáo dục thực chất:")
    
    add_body_p(doc, "Thứ nhất, tình trạng gian lận tiến độ học tập diễn ra phổ biến. Đa phần các web học tập thông thường chỉ ghi nhận sự kiện hoàn thành bài học khi con trỏ phát (playhead) chạm vào giây cuối cùng của video hoặc thông qua việc người học bấm nút 'Đánh dấu đã hoàn thành'. Học viên hoàn toàn có thể kéo thanh tua (seek bar) từ đầu tới cuối video trong 3 giây để được tính là hoàn thành khóa học mà không hề tiếp thu nội dung bài giảng.")
    add_body_p(doc, "Thứ hai, nguy cơ mất an toàn thông tin khi thực thi mã nguồn lập trình. Rất nhiều website đào tạo lập trình cho phép học viên chạy mã JavaScript trực tiếp trên trình duyệt bằng các hàm nguy hiểm như eval() hoặc new Function() mà không có cơ chế cô lập ngữ cảnh. Điều này cho phép mã độc đọc trộm localStorage, đánh cắp cookie chứa phiên đăng nhập JWT, hoặc gây treo toàn bộ ứng dụng thông qua các vòng lặp vô hạn (Infinite Loops).")
    add_body_p(doc, "Thứ ba, sự thiếu nhất quán trong quản lý tài chính và thương mại điện tử. Các giao dịch nạp tiền, sử dụng phiếu ưu đãi (voucher) và hoàn tiền (refund) thường bị xử lý phân tán, dễ dẫn đến tình trạng gian lận số dư, áp dụng mã giảm giá nhiều lần hoặc tranh chấp quyền học khi phát sinh hoàn tiền.")

    add_callout(doc,
                "Mục tiêu cốt lõi của Đồ án EduViet LMS: Xây dựng nền tảng học tập số chuyên sâu cho kỹ sư lập trình, "
                "kết hợp chặt chẽ giữa học lý thuyết qua video bài giảng chất lượng cao, thực hành code trên Web Worker Sandbox cô lập an toàn, "
                "làm bài kiểm tra trắc nghiệm được chấm điểm tự động tại server và kiểm soát thương mại minh bạch qua sổ cái bất biến (Wallet Ledger).",
                title="MỤC TIÊU CHIẾN LƯỢC DỰ ÁN EDUVIET")

    # 1.2 Đối tượng
    style_heading(doc.add_paragraph(), 2, "1.2 Phân tích đối tượng người dùng và khảo sát nhu cầu nghiệp vụ")
    add_body_p(doc, "Hệ thống EduViet LMS được mô hình hóa dựa trên ma trận phân quyền 3 vai trò chính (3-Role RBAC):")
    
    headers_roles = ["Nhóm người dùng", "Mô tả vai trò", "Nhu cầu nghiệp vụ then chốt", "Môi trường thao tác"]
    rows_roles = [
      ["Khách vãng lai (Guest)", "Người dùng chưa đăng ký hoặc chưa đăng nhập vào hệ thống.", "Tìm kiếm khóa học theo danh mục/giá/trình độ, xem chi tiết đề cương giáo trình, học thử bài giảng preview miễn phí.", "Trang chủ, Catalog, Chi tiết khóa học."],
      ["Học viên (Student)", "Người học đã xác thực tài khoản qua email/mật khẩu an toàn.", "Quản lý giỏ hàng, nạp tiền ví học tập, mua khóa học bằng ví, học video có lưu tiến độ chính xác, viết ghi chú theo mốc thời gian, nộp bài trắc nghiệm, luyện code trong Code Arena.", "Trang cá nhân, Phòng học (Learning Room), Luyện code, Lịch sử ví."],
      ["Giảng viên (Instructor)", "Chuyên gia phụ trách nội dung và chất lượng chuyên môn.", "Soạn thảo đề cương khóa học (Curriculum Editor: thêm chương, thêm bài học), quản lý tài liệu đính kèm, cập nhật link video học liệu, theo dõi danh sách học viên.", "Cổng giảng viên (Instructor Portal), Trình biên soạn giáo trình."],
      ["Quản trị viên (Admin)", "Người vận hành hệ thống cao nhất, đảm bảo tính tuân thủ.", "Kiểm duyệt và phê duyệt công bố khóa học (Moderation), đối soát yêu cầu nạp tiền, phê duyệt hoàn tiền và thu hồi quyền học, quản lý tài khoản người dùng và giám sát sổ cái tài chính.", "Bảng điều khiển quản trị (Admin Dashboard)."]
    ]
    add_styled_table(doc, headers_roles, rows_roles, [1.4, 1.8, 2.8, 1.2])

    # 1.3 36 User Stories
    style_heading(doc.add_paragraph(), 2, "1.3 Đặc tả chi tiết 36 User Stories và Tiêu chí nghiệm thu (Acceptance Criteria)")
    add_body_p(doc, "Toàn bộ chức năng của hệ thống EduViet được chia thành 36 User Stories chuẩn hóa (25 User Stories độ ưu tiên P1 bắt buộc hoàn thành và 11 User Stories độ ưu tiên P2 mở rộng), phân bổ đều trên 10 chủ đề nghiệp vụ:")

    headers_us = ["Mã US", "Độ ưu tiên", "Nội dung User Story", "Tiêu chí nghiệm thu (Acceptance Criteria)"]
    rows_us = [
      ["US01", "P1 (Bắt buộc)", "Đăng ký tài khoản người dùng mới.", "Validate định dạng email chuẩn, mật khẩu tối thiểu 6 ký tự, mã hóa bcrypt/argon2 an toàn, từ chối email trùng lặp (HTTP 409)."],
      ["US02", "P1 (Bắt buộc)", "Đăng nhập và quản lý phiên làm việc.", "Xác thực thông tin qua JWT Bearer token có hạn sử dụng (7 ngày), từ chối tài khoản bị khóa (is_active = false) với HTTP 401."],
      ["US03", "P1 (Bắt buộc)", "Đặt lại mật khẩu qua token xác minh.", "Tạo token ngẫu nhiên bảo mật cao có hạn 15 phút, từ chối token sai hoặc đã qua sử dụng, mật khẩu cũ hết hiệu lực ngay lập tức."],
      ["US04", "P1 (Bắt buộc)", "Cập nhật hồ sơ cá nhân và đổi mật khẩu.", "Xác minh chính xác mật khẩu hiện tại trước khi cập nhật mật khẩu mới, cập nhật họ tên và ảnh đại diện thành công."],
      ["US05", "P1 (Bắt buộc)", "Tìm kiếm, lọc và phân trang khóa học.", "Lọc đa tiêu chí theo từ khóa, danh mục, cấp độ, dải giá min-max; sắp xếp mới nhất/giá tăng/giá giảm; từ chối min_price > max_price (HTTP 400)."],
      ["US06", "P1 (Bắt buộc)", "Xem chi tiết khóa học và học thử miễn phí.", "Khách vãng lai xem được đề cương và bài học thử (is_preview = true); các bài học tính phí bị ẩn hoàn toàn URL học liệu."],
      ["US07", "P2 (Mở rộng)", "Lưu khóa học vào danh sách yêu thích.", "Lưu trữ khóa học quan tâm, hiển thị trạng thái yêu thích, xóa khỏi danh sách khi không còn nhu cầu."],
      ["US08", "P1 (Bắt buộc)", "Quản lý giỏ hàng học viên trên server.", "Thêm khóa học vào giỏ, từ chối khóa học đã sở hữu hoặc đã có trong giỏ (HTTP 409), tính toán tổng tiền chính xác."],
      ["US09", "P1 (Bắt buộc)", "Áp dụng phiếu ưu đãi giảm giá (Voucher).", "Kiểm tra hạn sử dụng, điều kiện đơn tối thiểu, giới hạn lượt dùng; tính mức giảm giá phần trăm hoặc số tiền cố định tại server."],
      ["US10", "P1 (Bắt buộc)", "Tạo yêu cầu nạp tiền vào ví điện tử.", "Tạo yêu cầu nạp với trạng thái PENDING, sinh mã giao dịch duy nhất, không tự ý cộng tiền cho đến khi được đối soát."],
      ["US11", "P1 (Bắt buộc)", "Thanh toán đơn hàng bằng số dư ví.", "Kiểm tra số dư ví >= giá trị đơn; trừ tiền ví, ghi nhận sổ cái PURCHASE, tạo đơn hàng PAID và cấp quyền Enrollment trong 1 transaction."],
      ["US12", "P1 (Bắt buộc)", "Yêu cầu hoàn tiền và hủy đơn hàng.", "Học viên gửi yêu cầu hoàn tiền cho đơn PAID kèm lý do; chuyển trạng thái yêu cầu sang PENDING chờ quản trị viên xem xét."],
      ["US13", "P1 (Bắt buộc)", "Xem danh sách Khóa học của tôi.", "Truy vấn danh sách các khóa học học viên đã mua kèm tỷ lệ phần trăm tiến độ hoàn thành và nút 'Học tiếp'."],
      ["US14", "P1 (Bắt buộc)", "Phòng học thông minh và chuyển bài chính xác.", "Player chuyển video đúng mốc thời gian start_seconds/end_seconds của bài giảng, hỗ trợ autoplay chuyển bài tự động."],
      ["US15", "P1 (Bắt buộc)", "Lưu tiến độ học tập qua thuật toán Interval Union.", "Tách biệt vị trí dừng xem và thời gian xem thực tế; hợp nhất các khoảng thời gian phát; chỉ công nhận hoàn thành khi xem đạt >= 90%."],
      ["US16", "P1 (Bắt buộc)", "Ghi chú bài học theo mốc thời gian.", "Tạo, sửa, xóa ghi chú cá nhân gắn với giây phát của video; nhấp vào ghi chú tự động tua trình phát đến đúng đoạn."],
      ["US17", "P1 (Bắt buộc)", "Tải tài liệu đính kèm bài giảng có bản quyền.", "Kiểm tra quyền truy cập (Enrollment) trước khi cấp link tải tài liệu PDF/Source Code, ngăn chặn tải trái phép."],
      ["US18", "P1 (Bắt buộc)", "Thực hành lập trình trong Code Arena an toàn.", "Thực thi mã nguồn JavaScript trong Web Worker Sandbox cô lập hoàn toàn khỏi DOM, tự động ngắt sau 2000ms timeout chống treo tab."],
      ["US19", "P1 (Bắt buộc)", "Làm bài trắc nghiệm Quiz chấm điểm tại server.", "Đề thi gửi về client không lộ đáp án đúng; server chấm điểm chuẩn hóa 0-100 theo trọng số câu hỏi; đánh giá đạt >= 80%."],
      ["US20", "P2 (Mở rộng)", "Nộp bài tập dự án thực hành cuối khóa.", "Học viên gửi đường dẫn kho lưu trữ mã nguồn GitHub và tài liệu mô tả để giảng viên đánh giá."],
      ["US21", "P2 (Mở rộng)", "Hỏi đáp thảo luận theo từng bài giảng.", "Học viên đặt câu hỏi gắn với bài học; giảng viên phụ trách trả lời; hiển thị trạng thái đã giải đáp."],
      ["US22", "P2 (Mở rộng)", "Đánh giá và nhận xét khóa học.", "Chỉ học viên đã mua và hoàn thành tối thiểu 30% khóa học mới được để lại đánh giá sao và nhận xét thực chất."],
      ["US23", "P2 (Mở rộng)", "Gửi yêu cầu hỗ trợ kỹ thuật (Support Ticket).", "Tạo ticket yêu cầu trợ giúp kèm phân loại tài khoản/thanh toán/bài giảng, theo dõi tiến độ xử lý."],
      ["US24", "P2 (Mở rộng)", "Cấp chứng nhận hoàn thành khóa học (Certificate).", "Tự động cấp mã chứng nhận số duy nhất có thể tra cứu công khai khi hoàn thành 100% bài học và đạt quiz."],
      ["US25", "P2 (Mở rộng)", "Thiết lập mục tiêu học tập hàng tuần.", "Cài đặt số giờ học mong muốn trong tuần và nhận thông báo nhắc nhở tiến độ định kỳ."],
      ["US26", "P2 (Mở rộng)", "Lộ trình học tập cá nhân hóa (Career Roadmap).", "Xem biểu đồ lộ trình nghề nghiệp từ Fresher lên Senior Fullstack tương ứng với các khóa học cần tích lũy."],
      ["US27", "P1 (Bắt buộc)", "Giảng viên quản lý khóa học của chính mình.", "CRUD khóa học thuộc quyền sở hữu của giảng viên; khóa mới tạo ở trạng thái Draft; không thể chỉnh sửa khóa của giảng viên khác."],
      ["US28", "P1 (Bắt buộc)", "Giảng viên biên soạn đề cương giáo trình.", "CRUD chương học (CourseModule) và bài học (Lesson); hỗ trợ sắp xếp lại thứ tự bài giảng qua reorder API."],
      ["US29", "P1 (Bắt buộc)", "Giảng viên tải lên học liệu video và tài liệu.", "Gắn liên kết video an toàn (YouTube, Cloudinary, MP4 stream); cấu hình thời lượng bài học và cờ học thử preview."],
      ["US30", "P2 (Mở rộng)", "Giảng viên chấm bài tập dự án của học viên.", "Xem danh sách bài nộp, đối chiếu barem rubric, chấm điểm và gửi phản hồi nhận xét chi tiết cho người học."],
      ["US31", "P1 (Bắt buộc)", "Quản trị viên quản lý tài khoản và phân quyền.", "Xem danh sách người dùng, kích hoạt/khóa tài khoản vi phạm, gán vai trò Admin/Instructor/Student."],
      ["US32", "P1 (Bắt buộc)", "Quản trị viên kiểm duyệt công bố khóa học.", "Xem các khóa học chờ duyệt; phê duyệt chuyển trạng thái sang Published hoặc từ chối kèm lý do phản hồi."],
      ["US33", "P1 (Bắt buộc)", "Quản trị viên đối soát tài chính và duyệt hoàn tiền.", "Kiểm tra yêu cầu nạp tiền và hoàn tiền; duyệt hoàn tiền tự động cộng số dư vào ví và thu hồi Enrollment."],
      ["US34", "P2 (Mở rộng)", "Quản trị viên xử lý báo cáo vi phạm nội dung.", "Tiếp nhận khiếu nại về bản quyền, ngôn từ thiếu văn hóa và thực hiện gỡ bỏ nội dung vi phạm."],
      ["US35", "P2 (Mở rộng)", "Bảng điều khiển phân tích số liệu và thống kê.", "Biểu đồ trực quan doanh thu tổng hợp, tỷ lệ hoàn thành khóa học, số lượng học viên mới theo thời gian."],
      ["US36", "P1 (Bắt buộc)", "Nhật ký kiểm toán hệ thống và sao lưu dữ liệu.", "Ghi nhật ký vết các tác vụ quan trọng (đổi mật khẩu, duyệt nạp tiền, hoàn tiền); quy trình sao lưu và khôi phục CSDL an toàn."]
    ]
    add_styled_table(doc, headers_us, rows_us, [0.8, 1.2, 2.4, 2.8])

    # 1.4 Phi chức năng
    style_heading(doc.add_paragraph(), 2, "1.4 Yêu cầu phi chức năng (NFR) và Ma trận truy vết yêu cầu (Traceability Matrix)")
    add_body_p(doc, "Bên cạnh các yêu cầu chức năng nghiệp vụ, hệ thống EduViet phải thỏa mãn các tiêu chuẩn kỹ thuật phi chức năng khắt khe nhằm đảm bảo độ tin cậy và khả năng vận hành bền bỉ:")
    
    headers_nfr = ["Mã NFR", "Tên yêu cầu", "Tiêu chuẩn kỹ thuật định lượng", "Giải pháp hiện thực hóa"]
    rows_nfr = [
      ["NFR01", "Hiệu năng phản hồi (Performance)", "Thời gian đáp ứng của API (p95) < 300ms đối với các tác vụ đọc dữ liệu catalog, tìm kiếm và phân trang.", "Sử dụng NestJS Fastify/Express engine tối ưu, đánh chỉ mục Index B-Tree tại MySQL trên các cột tìm kiếm và khóa ngoại."],
      ["NFR02", "Bảo mật định danh (Security)", "100% mật khẩu được băm bằng thuật toán an toàn; khóa JWT secret tối thiểu 16 ký tự; cấm hardcode credentials trong source code.", "Sử dụng module crypto băm khóa kèm Salt; cấu hình biến môi trường qua file .env; kiểm tra độ dài secret ngay khi ứng dụng khởi động."],
      ["NFR03", "Tính toàn vẹn dữ liệu (Data Integrity)", "Mọi giao dịch trừ tiền ví, tạo đơn hàng và cấp quyền học phải tuân thủ chuẩn ACID (Nguyên tử, Nhất quán, Cô lập, Bền vững).", "Sử dụng TypeORM Transaction Runner; nếu xảy ra lỗi ở bất kỳ bước nào, toàn bộ trạng thái số dư ví và đơn hàng sẽ rollback 100%."],
      ["NFR04", "Khả năng tương thích (Compatibility)", "Giao diện hiển thị sắc nét, responsive mượt mà từ màn hình Mobile (375px), Tablet (768px) đến Desktop (1920px).", "Bố cục lưới Bootstrap Grid System, breakpoint chuẩn, áp dụng phong cách thiết kế Booking.com với thanh tìm kiếm viền vàng linh hoạt."],
      ["NFR05", "Bảo vệ Sandbox (Isolation)", "Mã nguồn học viên không được phép làm treo trình duyệt hoặc truy xuất dữ liệu cá nhân trong ứng dụng.", "Môi trường thực thi Web Worker riêng biệt không có quyền truy cập DOM/Storage; cơ chế timeout tự động ngắt sau 2000ms."]
    ]
    add_styled_table(doc, headers_nfr, rows_nfr, [0.8, 1.6, 2.4, 2.4])

    doc.add_page_break()

# End of preliminary and Chapter 1
