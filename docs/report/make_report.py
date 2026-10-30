# -*- coding: utf-8 -*-
"""
TRÌNH TẠO BÁO CÁO HỌC THUẬT TOÀN DIỆN CHO ĐỒ ÁN EDUVIET LMS
Sinh viên thực hiện (Nhóm 6):
  1. Phù Ngọc Hào - MSSV: 23010384
  2. Nguyễn Quốc Khải - MSSV: 23010403 (Trưởng nhóm)
  3. Nguyễn Hoàng Sơn - MSSV: 23010878
Giảng viên hướng dẫn: TS. Nguyễn Lệ Thu
Học phần: Thiết kế Web Nâng cao - Trường Đại học Phenikaa
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def style_heading(p, level, text):
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.bold = True
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(13, 59, 102)
        p.paragraph_format.space_before = Pt(22)
        p.paragraph_format.space_after = Pt(8)
    elif level == 2:
        run.font.size = Pt(13.5)
        run.font.color.rgb = RGBColor(29, 53, 87)
        p.paragraph_format.space_before = Pt(15)
        p.paragraph_format.space_after = Pt(6)
    elif level == 3:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(30, 41, 59)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
    return run

def add_body_p(doc, text, bold_prefix=None, italic=False, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(space_after)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = "Times New Roman"
        r_bold.font.size = Pt(12)
        r_bold.bold = True
        r_bold.font.color.rgb = RGBColor(20, 20, 20)

    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.italic = italic
    run.font.color.rgb = RGBColor(30, 30, 30)
    return p

def add_callout(doc, text, title=None, border_color="FFB700"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, "F8F9FA")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:top w:val="none"/>'
                        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>'
                        f'<w:bottom w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)
    
    cp = cell.paragraphs[0]
    cp.paragraph_format.line_spacing = 1.25
    cp.paragraph_format.space_after = Pt(2)
    if title:
        rt = cp.add_run(f"📌 {title}\n")
        rt.bold = True
        rt.font.name = "Times New Roman"
        rt.font.size = Pt(11)
        rt.font.color.rgb = RGBColor(13, 59, 102)
    rc = cp.add_run(text)
    rc.font.name = "Times New Roman"
    rc.font.size = Pt(10.5)
    rc.font.italic = True
    rc.font.color.rgb = RGBColor(50, 50, 50)
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)

def add_styled_table(doc, headers, rows_data, col_widths=None):
    tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_shading(hdr_cells[i], "1D3557")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row in enumerate(rows_data):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_color = "F8F9FA" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_shading(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(30, 30, 30)

    if col_widths:
        for r in tbl.rows:
            for c_idx, w in enumerate(col_widths):
                r.cells[c_idx].width = Inches(w)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(6)

def add_code_block(doc, code_str, caption=None):
    if caption:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(6)
        cp.paragraph_format.space_after = Pt(2)
        cr = cp.add_run(f"Đoạn mã: {caption}")
        cr.bold = True
        cr.font.name = "Times New Roman"
        cr.font.size = Pt(10.5)
        cr.font.color.rgb = RGBColor(50, 50, 50)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, "F1F5F9")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="0D3B66"/>'
                        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(code_str)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(15, 23, 42)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)

def setup_page_layout(doc):
    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(0.98)
        section.right_margin = Inches(0.79)

        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("TRƯỜNG ĐẠI HỌC PHENIKAA — KHOA CNTT | BÁO CÁO ĐỒ ÁN EDUVIET LMS")
        hrun.font.name = "Times New Roman"
        hrun.font.size = Pt(8.5)
        hrun.font.italic = True
        hrun.font.color.rgb = RGBColor(120, 120, 120)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun = fp.add_run("Học phần: Thiết kế Web Nâng cao — GVHD: TS. Nguyễn Lệ Thu")
        frun.font.name = "Times New Roman"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(120, 120, 120)
